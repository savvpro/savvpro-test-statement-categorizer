"""Bounded statement extraction and explicit column mapping. Never discard bad rows."""

import csv
import io
import re
import zipfile
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
ALIASES = {
    "date": ["date", "transaction date", "posting date", "posted date"],
    "description": ["description", "transaction description", "details", "narration", "merchant"],
    "amount": ["amount", "transaction amount", "posted amount"],
    "debit": ["debit", "debit amount", "withdrawal", "withdrawals"],
    "credit": ["credit", "credit amount", "deposit", "deposits"],
    "currency": ["currency", "ccy"],
    "original_amount": ["original amount", "foreign amount"],
    "original_currency": ["original currency", "foreign currency"],
}


class ImportProblem(ValueError):
    pass


def extract(filename: str, content: bytes) -> dict:
    if not content or len(content) > MAX_BYTES:
        raise ImportProblem("Choose a non-empty statement smaller than 5 MB.")
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = content.decode("cp1252")
            except UnicodeDecodeError as exc:
                raise ImportProblem("CSV encoding is unsupported. Export it as UTF-8 CSV.") from exc
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        try:
            rows = list(csv.reader(io.StringIO(text), dialect=dialect))
        except csv.Error as exc:
            raise ImportProblem("Malformed CSV. Check quoting and column separators.") from exc
    elif suffix == ".xlsx":
        from openpyxl import load_workbook

        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024:
                    raise ImportProblem("Workbook expands beyond the 40 MB safety limit.")
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            if len(book.sheetnames) != 1:
                book.close()
                raise ImportProblem(
                    "Use a workbook with one statement sheet, or export the intended sheet as CSV."
                )
            rows = []
            for row in book.active.iter_rows(values_only=True):
                rows.append(
                    [str(v.date() if isinstance(v, datetime) else v) if v is not None else "" for v in row]
                )
                if len(rows) > MAX_ROWS + 1:
                    break
            book.close()
        except ImportProblem:
            raise
        except Exception as exc:
            raise ImportProblem("Unable to read this workbook. Use an unlocked .xlsx or CSV file.") from exc
    elif suffix == ".pdf":
        import pdfplumber

        try:
            with pdfplumber.open(io.BytesIO(content)) as document:
                if len(document.pages) > 30:
                    raise ImportProblem("PDF limit is 30 pages.")
                text = "\n".join(page.extract_text() or "" for page in document.pages)
                tables = [table for page in document.pages for table in page.extract_tables()]
        except ImportProblem:
            raise
        except Exception as exc:
            raise ImportProblem("Unable to read PDF. Upload an unlocked text PDF or a CSV export.") from exc
        # A deliberately defined format, not a claim of universal bank PDF support.
        if "SortedHisab Synthetic Statement v1" not in text:
            raise ImportProblem(
                "PDF layout not supported. Use the supplied synthetic v1 PDF or export your statement as CSV/XLSX. Scans are not supported."
            )
        rows = [["Date", "Description", "Debit", "Credit", "Currency"]]
        for table in tables:
            if table and table[0] == rows[0]:
                for cells in table[1:]:
                    if len(cells) != 5:
                        raise ImportProblem(
                            "A PDF transaction line is incomplete; import stopped to prevent missing rows."
                        )
                    rows.append([cell or "" for cell in cells])
        expected = re.search(r"Transaction count: (\d+)", text)
        if not expected or int(expected[1]) != len(rows) - 1:
            raise ImportProblem("PDF row count does not reconcile. Use CSV instead.")
    else:
        raise ImportProblem("Supported files: CSV, XLSX, and the synthetic v1 text PDF.")
    rows = [row for row in rows if any(str(cell).strip() for cell in row)]
    if len(rows) < 2:
        raise ImportProblem("Statement must contain a header and at least one transaction.")
    if len(rows) > MAX_ROWS + 1:
        raise ImportProblem(f"Statement exceeds {MAX_ROWS:,} transaction rows.")
    headers = [str(value).strip() for value in rows[0]]
    if len(headers) > 40 or len(set(headers)) != len(headers) or any(not h for h in headers):
        raise ImportProblem("Use at most 40 uniquely named, non-empty column headers.")
    mapping = {}
    for key, aliases in ALIASES.items():
        mapping[key] = next((h for h in headers if h.lower() in aliases), "")
    return {"headers": headers, "rows": rows[1:], "mapping": mapping, "row_count": len(rows) - 1}


def money(value: str) -> int:
    text = str(value).strip().upper()
    if not text:
        raise ImportProblem("Amount is empty.")
    negative = text.startswith("(") and text.endswith(")")
    text = re.sub(r"^(PKR|USD|EUR|GBP|AED|RS\.?)[ ]*", "", text)
    text = text.replace(",", "").strip("() ")
    # DR/CR must be resolved by explicit mapping/sign convention, never guessed here.
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ImportProblem("Invalid amount; use numbers with a decimal point and no DR/CR suffix.") from exc
    if not number.is_finite() or abs(number) > Decimal("100000000000"):
        raise ImportProblem("Amount is not finite or exceeds the supported range.")
    if number != number.quantize(Decimal("0.01")):
        raise ImportProblem("Use at most two decimal places.")
    return int(number * 100) * (-1 if negative else 1)


def parse_rows(extracted: dict, options: dict) -> tuple[list, list]:
    headers, mapping = extracted["headers"], options["mapping"]
    if not mapping.get("date") or not mapping.get("description"):
        raise ImportProblem("Map a date and description column.")
    if not mapping.get("amount") and not (mapping.get("debit") and mapping.get("credit")):
        raise ImportProblem("Map an amount column or both debit and credit columns.")
    if mapping.get("amount") and (mapping.get("debit") or mapping.get("credit")):
        raise ImportProblem("Choose either one amount column or debit/credit columns, not both.")
    used = [v for v in mapping.values() if v]
    if any(v not in headers for v in used) or len(used) != len(set(used)):
        raise ImportProblem("Mapped columns must exist and cannot be reused.")
    transactions, errors = [], []
    for index, cells in enumerate(extracted["rows"], start=2):
        row = dict(zip(headers, cells))

        def get(key):
            return str(row.get(mapping.get(key, ""), "")).strip()

        try:
            if len(cells) != len(headers):
                raise ImportProblem("Row has a different column count than the header.")
            raw_date = get("date")
            parsed_date = None
            formats = ["%Y-%m-%d", "%d %b %Y", "%d-%b-%Y"]
            formats += (
                ["%d/%m/%Y", "%d-%m-%Y"] if options["date_order"] == "day_first" else ["%m/%d/%Y", "%m-%d-%Y"]
            )
            for fmt in formats:
                try:
                    parsed_date = datetime.strptime(raw_date, fmt).date().isoformat()
                    break
                except ValueError:
                    pass
            if not parsed_date:
                raise ImportProblem("Invalid date for the selected date order.")
            description = get("description")
            if not description or len(description) > 500:
                raise ImportProblem("Description must contain 1–500 characters.")
            if mapping.get("amount"):
                amount = money(get("amount"))
                if options["amount_direction"] == "in_positive":
                    amount *= -1
            else:
                debit, credit = money(get("debit") or "0"), money(get("credit") or "0")
                if debit < 0 or credit < 0 or (debit and credit):
                    raise ImportProblem(
                        "Debit/credit cells must be nonnegative with only one populated amount."
                    )
                amount = debit - credit
            if amount == 0:
                raise ImportProblem("Zero amount: check whether this is a transaction or statement metadata.")
            currency = (get("currency") or options["currency"]).upper()
            if currency not in {"PKR", "USD", "EUR", "GBP", "AED"}:
                raise ImportProblem("Supported currencies: PKR, USD, EUR, GBP, AED.")
            original_amount, original_currency = get("original_amount"), get("original_currency").upper()
            if bool(original_amount) != bool(original_currency):
                raise ImportProblem("Original amount and original currency must be supplied together.")
            if original_currency and original_currency not in {"PKR", "USD", "EUR", "GBP", "AED"}:
                raise ImportProblem("Unsupported original currency.")
            transactions.append(
                {
                    "date": parsed_date,
                    "description": description,
                    "amount_minor": amount,
                    "currency": currency,
                    "source_row": index,
                    "original_amount_minor": money(original_amount) if original_amount else None,
                    "original_currency": original_currency or None,
                    "raw": row,
                }
            )
        except ImportProblem as exc:
            errors.append({"row": index, "message": str(exc), "raw": row})
    return transactions, errors
