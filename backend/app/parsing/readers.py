"""Turn raw upload bytes into a grid of string cells, keeping the source line numbers."""

import codecs
import csv
import datetime as dt
import io
from collections import Counter
from dataclasses import dataclass

from app.parsing.errors import EmptyFileError, UnsupportedFileError

ZIP_MAGIC = b"PK\x03\x04"
_OLE_MAGIC = b"\xd0\xcf\x11\xe0"  # legacy .xls / .doc
_DELIMITERS = (",", ";", "\t", "|")
_SNIFF_LINES = 30


@dataclass(frozen=True, slots=True)
class SourceRow:
    line_no: int  # 1-based line (csv) or row number (xlsx)
    cells: list[str]


@dataclass(frozen=True, slots=True)
class Table:
    file_type: str
    encoding: str | None
    delimiter: str | None
    rows: list[SourceRow]


def read_table(content: bytes, filename: str) -> Table:
    if not content.strip():
        raise EmptyFileError(f"{filename} is empty")
    if content.startswith(ZIP_MAGIC):
        return _read_xlsx(content, filename)
    if content.startswith(_OLE_MAGIC):
        raise UnsupportedFileError(
            f"{filename}: legacy .xls is not supported; save as .xlsx or .csv"
        )
    return _read_csv(content, filename)


def _read_csv(content: bytes, filename: str) -> Table:
    text, encoding = _decode(content, filename)
    if not text.strip():
        raise EmptyFileError(f"{filename} is empty")

    delimiter = sniff_delimiter(text)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    rows: list[SourceRow] = []
    next_line = 1
    try:
        for cells in reader:
            rows.append(SourceRow(next_line, [cell.strip() for cell in cells]))
            next_line = reader.line_num + 1
    except csv.Error as exc:
        raise UnsupportedFileError(
            f"{filename}: malformed CSV near line {next_line}: {exc}"
        ) from exc
    return Table("csv", encoding, delimiter, rows)


def _decode(content: bytes, filename: str) -> tuple[str, str]:
    # A BOM settles it; otherwise try strict UTF-8 first, then the common Windows export encoding.
    candidates = ["utf-8-sig"] if content.startswith(codecs.BOM_UTF8) else ["utf-8", "cp1252"]

    for encoding in candidates:
        try:
            text = content.decode(encoding)
        except UnicodeDecodeError:
            continue
        if _looks_binary(text):
            break
        return text, encoding
    raise UnsupportedFileError(f"{filename} is not a CSV/text file or an .xlsx workbook")


def _looks_binary(text: str) -> bool:
    sample = text[:8192]
    if "\x00" in sample:
        return True
    control = sum(1 for ch in sample if ord(ch) < 32 and ch not in "\t\r\n")
    return control > len(sample) * 0.01


def sniff_delimiter(text: str) -> str:
    """Pick the delimiter that gives the most rows with the same (>1) number of fields.

    Built on `csv.reader` so quoted fields like "1,234.56" do not confuse the count, and robust
    to junk lines above the header (they simply do not vote).
    """
    lines = [line for line in text.splitlines() if line.strip()][:_SNIFF_LINES]
    best, best_score = ",", (0, 0)
    for delimiter in _DELIMITERS:
        widths = [len(row) for row in csv.reader(lines, delimiter=delimiter) if len(row) > 1]
        if not widths:
            continue
        width, votes = Counter(widths).most_common(1)[0]
        score = (votes, width)
        if score > best_score:
            best, best_score = delimiter, score
    return best


def _read_xlsx(content: bytes, filename: str) -> Table:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for non-workbook zips
        raise UnsupportedFileError(f"{filename} is not a valid .xlsx workbook") from exc

    try:
        if not workbook.worksheets:
            raise EmptyFileError(f"{filename} has no worksheets")
        sheet = workbook.worksheets[0]
        rows = [
            SourceRow(line_no, [_cell_text(value) for value in values])
            for line_no, values in enumerate(sheet.iter_rows(values_only=True), start=1)
        ]
    finally:
        workbook.close()
    return Table("xlsx", None, None, rows)


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float):
        # repr gives the shortest string that round-trips, e.g. 45.1 -> "45.1"; no float math.
        return repr(value)
    return str(value).strip()
