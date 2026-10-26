"""Excel workbooks (.xlsx)."""

from __future__ import annotations

import io
import zipfile
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from openpyxl import load_workbook

from app.domain.models import ParsedStatement, SourceFormat
from app.errors import FileProblem
from app.parsing.tabular import find_columns, grid_to_statement

_MAX_UNCOMPRESSED = 60 * 1024 * 1024
_MAX_SHEETS = 10


def _check_archive(data: bytes) -> None:
    """Refuse zip bombs before handing the workbook to the XML parser."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            total = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise FileProblem("The Excel file is damaged and could not be opened.") from exc
    if total > _MAX_UNCOMPRESSED:
        raise FileProblem(
            "The workbook expands to an unreasonable size and was rejected.",
            code="file_too_large",
            status_code=413,
        )


def _to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return f"{Decimal(repr(value)):f}"
    return str(value)


def read_xlsx(data: bytes, *, max_rows: int) -> ParsedStatement:
    _check_archive(data)
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a wide variety of errors on bad input
        raise FileProblem("The Excel file is damaged and could not be opened.") from exc

    try:
        last_error: FileProblem | None = None
        for sheet in workbook.worksheets[:_MAX_SHEETS]:
            grid: list[list[str]] = []
            line_numbers: list[int] = []
            for number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                cells = [_to_text(v) for v in row]
                if any(c.strip() for c in cells):
                    grid.append(cells)
                    line_numbers.append(number)
                if len(grid) > max_rows + 60:
                    break
            if find_columns(grid) is None:
                continue
            try:
                parsed = grid_to_statement(
                    grid, SourceFormat.XLSX, max_rows=max_rows, line_numbers=line_numbers
                )
            except FileProblem as exc:
                last_error = exc
                continue
            if len(workbook.worksheets) > 1:
                parsed.warnings.append(f"Read the sheet named {sheet.title!r}.")
            return parsed
        if last_error is not None:
            raise last_error
    finally:
        workbook.close()
    raise FileProblem(
        "None of the sheets has a header row with date, description and amount columns.",
        code="missing_columns",
    )
